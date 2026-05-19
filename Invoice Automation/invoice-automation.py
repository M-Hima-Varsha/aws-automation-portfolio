import os
import re
import sys
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor
from copy import copy

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.dont_write_bytecode = True

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PDF_DIR = os.path.join(SCRIPT_DIR, "invoices")
REPORT_FILE_PREFIX = "Invoice_Report_Dev&Prod"
OUTPUT_FILE = None
MAX_WORKERS = 5

ACCOUNT_PATTERN = re.compile(r"^(?P<name>.+?) \((?P<id>\d{12})\) USD (?P<amount>-?[\d,]+\.\d+)$")
SERVICE_PATTERN = re.compile(r"^(?P<name>.+?) USD (?P<amount>-?[\d,]+\.\d+)$")
INVOICE_PATTERN = re.compile(r"Invoice Number:\s*([A-Z0-9-]+)")
SUMMARY_CHARGES_PATTERN = re.compile(r"^Charges USD (?P<amount>-?[\d,]+\.\d+)$")
LATE_FEE_PATTERN = re.compile(r"^(?P<name>OCBLateFee) USD (?P<amount>-?[\d,]+\.\d+)$", re.IGNORECASE)
BILLING_PERIOD_PATTERN = re.compile(
    r"This invoice is for the billing period\s+"
    r"(?P<start_month>[A-Za-z]+)\s+\d+\s*-\s*(?P<end_month>[A-Za-z]+)\s+\d+\s*,\s*(?P<year>\d{4})",
    re.IGNORECASE,
)

IGNORED_DETAIL_PREFIXES = (
    "Charges USD",
    "Credits USD",
    "Discount",
    "Estimated US sales tax",
    "Account ",
    "For line item details",
    "* May include",
)


def parse_amount(amount_text):
    return float(amount_text.replace(",", ""))


def normalize_text(value):
    text = str(value or "").strip().lower()
    return " ".join(text.split())


def normalize_account_lookup(value):
    return normalize_text(value).rstrip(".")


def extract_billing_period_label_from_text(text):
    match = BILLING_PERIOD_PATTERN.search(text or "")
    if not match:
        return None, None
    month_name = match.group("start_month").strip()
    year = match.group("year").strip()
    return month_name, f"{month_name}-{year}"


def get_billing_period_from_pdf(pdf_path):
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages[:2]:
                text = page.extract_text() or ""
                month_name, label = extract_billing_period_label_from_text(text)
                if label:
                    return month_name, label
    except Exception as exc:
        print(f"Error reading billing period from {pdf_path}: {exc}")

    return None, None


def get_expected_previous_billing_period():
    now = datetime.now()
    year = now.year
    month = now.month - 1

    if month == 0:
        month = 12
        year -= 1

    month_name = datetime(year, month, 1).strftime("%B")
    return month_name, f"{month_name}-{year}"


def determine_billing_period(pdf_dir):
    pdf_files = sorted(
        os.path.join(pdf_dir, file_name)
        for file_name in os.listdir(pdf_dir)
        if file_name.lower().endswith(".pdf")
    )

    period_map = {}

    for pdf_file in pdf_files:
        current_month_name, current_label = get_billing_period_from_pdf(pdf_file)
        if current_label:
            period_map[current_label] = current_month_name

    if not period_map:
        return None, None, []

    sorted_labels = sorted(period_map)
    if len(sorted_labels) == 1:
        selected_label = sorted_labels[0]
        return period_map[selected_label], selected_label, sorted_labels

    print("Warning: multiple billing periods found in invoices:")
    for label in sorted_labels:
        print(f" - {label}")

    return None, None, sorted_labels


def validate_invoices_match_previous_month(pdf_dir):
    pdf_files = sorted(
        os.path.join(pdf_dir, file_name)
        for file_name in os.listdir(pdf_dir)
        if file_name.lower().endswith(".pdf")
    )

    if not pdf_files:
        raise SystemExit(f"No PDF invoices found in '{pdf_dir}'.")

    expected_month_name, expected_label = get_expected_previous_billing_period()
    mismatched_invoices = []
    unreadable_invoices = []

    for pdf_file in pdf_files:
        _, current_label = get_billing_period_from_pdf(pdf_file)
        if not current_label:
            unreadable_invoices.append(os.path.basename(pdf_file))
            continue

        if current_label != expected_label:
            mismatched_invoices.append(
                {
                    "file": os.path.basename(pdf_file),
                    "billing_period": current_label,
                }
            )

    if unreadable_invoices or mismatched_invoices:
        print(
            f"Validation failed. All invoices must have billing period '{expected_label}' "
            f"based on the local system month."
        )

        for file_name in unreadable_invoices:
            print(f" - {file_name}: billing period not found")

        for invoice in mismatched_invoices:
            print(f" - {invoice['file']}: found '{invoice['billing_period']}'")

        raise SystemExit("Stopping execution because invoice billing period validation failed.")

    return expected_month_name, expected_label


def prompt_for_billing_period(period_labels):
    print("Select the billing month for the output file:")
    for index, label in enumerate(period_labels, start=1):
        print(f"{index}. {label}")

    while True:
        selected_value = input("Enter option number: ").strip()
        if selected_value.isdigit():
            selected_index = int(selected_value)
            if 1 <= selected_index <= len(period_labels):
                selected_label = period_labels[selected_index - 1]
                selected_month = selected_label.rsplit("-", 1)[0]
                return selected_month, selected_label
        print("Invalid selection. Please enter one of the listed option numbers.")


def build_output_file_name(billing_period_label=None):
    _, billing_period_label = get_expected_previous_billing_period()
    if not billing_period_label:
        if OUTPUT_FILE:
            return OUTPUT_FILE if os.path.isabs(OUTPUT_FILE) else os.path.join(SCRIPT_DIR, OUTPUT_FILE)
        return os.path.join(SCRIPT_DIR, f"{REPORT_FILE_PREFIX}.xlsx")

    updated_name = f"{REPORT_FILE_PREFIX}-{billing_period_label}"

    return os.path.join(SCRIPT_DIR, f"{updated_name}.xlsx")


def get_available_output_path(output_file):
    if not os.path.exists(output_file):
        return output_file

    file_name, extension = os.path.splitext(output_file)
    index = 1
    while True:
        candidate = f"{file_name}-{index}{extension}"
        if not os.path.exists(candidate):
            return candidate
        index += 1


def new_account(invoice_number, invoice_file, account_name, account_id, total_charge):
    return {
        "invoice_number": invoice_number,
        "invoice_file": invoice_file,
        "account_name": account_name,
        "account_id": account_id,
        "account_total": total_charge,
        "summary_charges": None,
        "services": [],
        "marketplace": [],
        "late_fees": [],
    }


def extract_invoice_header_details(text, default_invoice_number):
    invoice_number = default_invoice_number
    bill_to_name = None
    account_id = None
    lines = [line.strip() for line in (text or "").splitlines() if line.strip()]

    invoice_match = INVOICE_PATTERN.search(text or "")
    if invoice_match:
        invoice_number = invoice_match.group(1)

    for index, line in enumerate(lines):
        if line.lower() == "account number:":
            for next_line in lines[index + 1 :]:
                if re.fullmatch(r"\d{12}", next_line):
                    account_id = next_line
                    break
            continue

        if line.startswith("Bill to Address:"):
            for next_line in lines[index + 1 :]:
                if next_line.startswith("Please include this invoice number"):
                    break
                bill_to_name = next_line
                break

    return invoice_number, bill_to_name, account_id


def format_late_fee_remark(fee_name):
    if normalize_text(fee_name) == "ocblatefee":
        return "OBC Late Fee"
    return fee_name


def should_skip_detail_line(line):
    if not line or "USD" not in line:
        return True
    return any(line.startswith(prefix) for prefix in IGNORED_DETAIL_PREFIXES)


def parse_invoice(pdf_path):
    invoice_results = []
    current_section = None
    current_account = None
    invoice_number = os.path.splitext(os.path.basename(pdf_path))[0]
    bill_to_name = None
    bill_to_account_id = None
    invoice_level_late_fees = []

    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page_index, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if not text:
                    continue

                if page_index == 0:
                    invoice_number, bill_to_name, bill_to_account_id = extract_invoice_header_details(
                        text,
                        invoice_number,
                    )
                elif invoice_number.startswith("Invoice_"):
                    invoice_match = INVOICE_PATTERN.search(text)
                    if invoice_match:
                        invoice_number = invoice_match.group(1)

                for raw_line in text.splitlines():
                    line = raw_line.strip()
                    if not line:
                        continue

                    if line == "Summary" or "Summary for Linked Account" in line:
                        current_section = "summary"
                        continue

                    if line == "Detail" or "Detail for Linked Account" in line:
                        current_section = "detail"
                        continue

                    account_match = ACCOUNT_PATTERN.match(line)
                    if account_match:
                        if current_account:
                            invoice_results.append(current_account)

                        account_name = account_match.group("name").strip()
                        account_id = account_match.group("id").strip()
                        total_charge = parse_amount(account_match.group("amount"))

                        current_account = new_account(
                            invoice_number,
                            os.path.basename(pdf_path),
                            account_name,
                            account_id,
                            total_charge,
                        )
                        continue

                    if current_section == "summary" and current_account:
                        summary_match = SUMMARY_CHARGES_PATTERN.match(line)
                        if summary_match and current_account["summary_charges"] is None:
                            current_account["summary_charges"] = parse_amount(summary_match.group("amount"))
                        continue

                    if current_section == "detail" and not current_account:
                        late_fee_match = LATE_FEE_PATTERN.match(line)
                        if late_fee_match:
                            invoice_level_late_fees.append(
                                {
                                    "name": late_fee_match.group("name").strip(),
                                    "charge": parse_amount(late_fee_match.group("amount")),
                                }
                            )
                        continue

                    if current_section != "detail" or not current_account or should_skip_detail_line(line):
                        continue

                    service_match = SERVICE_PATTERN.match(line)
                    if not service_match:
                        continue

                    service_name = service_match.group("name").strip()
                    service_charge = parse_amount(service_match.group("amount"))
                    service_entry = {"name": service_name, "charge": service_charge}

                    if " sold by " in service_name.lower():
                        current_account["marketplace"].append(service_entry)
                    else:
                        current_account["services"].append(service_entry)

        if current_account:
            invoice_results.append(current_account)

        if invoice_level_late_fees and bill_to_name:
            late_fee_account = new_account(
                invoice_number,
                os.path.basename(pdf_path),
                bill_to_name,
                bill_to_account_id or "",
                round(sum(item["charge"] for item in invoice_level_late_fees), 2),
            )
            late_fee_account["late_fees"] = invoice_level_late_fees
            invoice_results.append(late_fee_account)

    except Exception as exc:
        print(f"Error processing {pdf_path}: {exc}")

    return invoice_results


def process_all_invoices(pdf_dir, selected_billing_period_label=None):
    pdf_files = [
        os.path.join(pdf_dir, file_name)
        for file_name in os.listdir(pdf_dir)
        if file_name.lower().endswith(".pdf")
    ]

    if selected_billing_period_label:
        filtered_pdf_files = []
        for pdf_file in pdf_files:
            _, current_label = get_billing_period_from_pdf(pdf_file)
            if current_label == selected_billing_period_label:
                filtered_pdf_files.append(pdf_file)
        pdf_files = filtered_pdf_files

    all_results = []
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        futures = [executor.submit(parse_invoice, pdf_file) for pdf_file in pdf_files]
        for future in futures:
            result = future.result()
            if result:
                all_results.extend(result)

    all_results = resolve_late_fee_account_names(all_results)
    return consolidate_results(all_results)


def consolidate_results(data):
    consolidated = {}

    for item in data:
        key = (
            item["invoice_file"],
            item["invoice_number"],
            item["account_name"],
            item["account_id"],
        )

        if key not in consolidated:
            consolidated[key] = {
                "invoice_file": item["invoice_file"],
                "invoice_number": item["invoice_number"],
                "account_name": item["account_name"],
                "account_id": item["account_id"],
                "account_total": item["account_total"],
                "summary_charges": item["summary_charges"],
                "services": [],
                "marketplace": [],
                "late_fees": [],
            }
        else:
            if consolidated[key]["account_total"] is None:
                consolidated[key]["account_total"] = item["account_total"]
            if consolidated[key]["summary_charges"] is None:
                consolidated[key]["summary_charges"] = item["summary_charges"]

        consolidated[key]["services"].extend(item["services"])
        consolidated[key]["marketplace"].extend(item["marketplace"])
        consolidated[key]["late_fees"].extend(item.get("late_fees", []))

    for item in consolidated.values():
        seen_services = set()
        deduped_services = []
        for service in item["services"]:
            service_key = (service["name"], service["charge"])
            if service_key not in seen_services:
                seen_services.add(service_key)
                deduped_services.append(service)
        item["services"] = deduped_services

        seen_marketplace = set()
        deduped_marketplace = []
        for marketplace in item["marketplace"]:
            marketplace_key = (marketplace["name"], marketplace["charge"])
            if marketplace_key not in seen_marketplace:
                seen_marketplace.add(marketplace_key)
                deduped_marketplace.append(marketplace)
        item["marketplace"] = deduped_marketplace

        seen_late_fees = set()
        deduped_late_fees = []
        for late_fee in item["late_fees"]:
            late_fee_key = (late_fee["name"], late_fee["charge"])
            if late_fee_key not in seen_late_fees:
                seen_late_fees.add(late_fee_key)
                deduped_late_fees.append(late_fee)
        item["late_fees"] = deduped_late_fees

    return sorted(
        consolidated.values(),
        key=lambda item: (item["invoice_number"], item["account_name"], item["invoice_file"]),
    )


def resolve_late_fee_account_names(data):
    account_id_to_name = {}

    for item in data:
        account_id = str(item.get("account_id") or "").strip()
        has_non_late_fee_entries = (
            item.get("summary_charges") is not None
            or bool(item.get("services"))
            or bool(item.get("marketplace"))
        )
        if not account_id or not has_non_late_fee_entries:
            continue

        account_id_to_name.setdefault(account_id, set()).add(item["account_name"])

    resolved_account_names = {
        account_id: next(iter(account_names))
        for account_id, account_names in account_id_to_name.items()
        if len(account_names) == 1
    }

    for item in data:
        account_id = str(item.get("account_id") or "").strip()
        has_only_late_fees = (
            bool(item.get("late_fees"))
            and item.get("summary_charges") is None
            and not item.get("services")
            and not item.get("marketplace")
        )
        if not account_id or not has_only_late_fees:
            continue

        resolved_name = resolved_account_names.get(account_id)
        if resolved_name:
            item["account_name"] = resolved_name

    return data


def get_service_total(account):
    total_service_charge = account["summary_charges"]
    if total_service_charge is None:
        total_service_charge = account["account_total"]
    if total_service_charge is None:
        total_service_charge = round(sum(service["charge"] for service in account["services"]), 2)
    return round(total_service_charge, 2)


def build_account_entries(data):
    account_entries = {}

    for account in data:
        account_key = normalize_account_lookup(account["account_name"])
        marketplace_total = round(sum(item["charge"] for item in account["marketplace"]), 2)
        entry_rows = []
        has_service_entries = bool(account["services"] or account["marketplace"]) or account["summary_charges"] is not None

        if has_service_entries:
            service_total = get_service_total(account)
            aws_service_total = round(service_total - marketplace_total, 2)
            if account["marketplace"]:
                aws_service_amount = max(aws_service_total, 0)
                if aws_service_amount > 0:
                    entry_rows.append(
                        {
                            "invoice_number": account["invoice_number"],
                            "amount": round(aws_service_amount, 2),
                            "remark": "AWS Services Cost",
                        }
                    )
            else:
                entry_rows.append(
                    {
                        "invoice_number": account["invoice_number"],
                        "amount": service_total,
                        "remark": "AWS Services Cost",
                    }
                )

        for marketplace in account["marketplace"]:
            entry_rows.append(
                {
                    "invoice_number": account["invoice_number"],
                    "amount": round(marketplace["charge"], 2),
                    "remark": f"AWS Market Place Charges - {marketplace['name']}",
                }
            )

        for late_fee in account.get("late_fees", []):
            entry_rows.append(
                {
                    "invoice_number": account["invoice_number"],
                    "amount": round(late_fee["charge"], 2),
                    "remark": format_late_fee_remark(late_fee["name"]),
                }
            )

        if not entry_rows:
            continue

        account_entries.setdefault(account_key, []).extend(entry_rows)

    for account_key, entries in account_entries.items():
        entries.sort(key=lambda item: (item["invoice_number"], item["remark"], item["amount"]))
        account_entries[account_key] = deduplicate_entries(entries)

    return account_entries


def is_title_or_header(cell_value):
    text = normalize_text(cell_value)
    return text.startswith("invoice") or text == "account"


def row_has_account_template_values(account_value, env_value):
    account_text = normalize_text(account_value)
    env_text = normalize_text(env_value)
    return (
        bool(account_text)
        and bool(env_text)
        and account_text != "account"
        and env_text != "environment type"
        and not account_text.startswith("invoice")
    )


def collect_template_rows(workbook):
    template_rows = []

    for worksheet in workbook.worksheets:
        for row_index in range(1, worksheet.max_row + 1):
            account_value = worksheet.cell(row=row_index, column=1).value
            env_value = worksheet.cell(row=row_index, column=2).value

            if normalize_text(account_value) == "account":
                continue

            if is_title_or_header(account_value):
                continue

            if not row_has_account_template_values(account_value, env_value):
                continue

            template_rows.append(
                {
                    "sheet": worksheet.title,
                    "row": row_index,
                    "block_end_row": get_account_block_end(worksheet, row_index),
                    "account_name": str(account_value).strip(),
                }
            )

    return template_rows


def get_account_block_end(worksheet, start_row):
    end_row = start_row

    for row_index in range(start_row + 1, worksheet.max_row + 1):
        account_value = worksheet.cell(row=row_index, column=1).value
        env_value = worksheet.cell(row=row_index, column=2).value
        month_value = worksheet.cell(row=row_index, column=3).value
        invoice_value = worksheet.cell(row=row_index, column=4).value
        amount_value = worksheet.cell(row=row_index, column=5).value
        remark_value = worksheet.cell(row=row_index, column=7).value

        if is_title_or_header(account_value):
            break

        if row_has_account_template_values(account_value, env_value):
            break

        if (
            account_value is None
            and env_value is None
            and month_value is None
            and invoice_value is None
            and amount_value is None
            and remark_value is None
        ):
            break

        end_row = row_index

    return end_row


def deduplicate_entries(entries):
    seen_entries = set()
    unique_entries = []

    for entry in entries:
        entry_key = (
            str(entry.get("invoice_number") or "").strip(),
            round(float(entry.get("amount") or 0), 2),
            str(entry.get("remark") or "").strip(),
        )
        if entry_key in seen_entries:
            continue
        seen_entries.add(entry_key)
        unique_entries.append(
            {
                "invoice_number": entry_key[0],
                "amount": entry_key[1],
                "remark": entry_key[2],
            }
        )

    unique_entries.sort(key=lambda item: (item["invoice_number"], item["remark"], item["amount"]))
    return unique_entries


def collect_existing_account_entries(workbook):
    existing_entries = {}
    template_rows = collect_template_rows(workbook)

    for template_row in template_rows:
        worksheet = workbook[template_row["sheet"]]
        account_key = normalize_account_lookup(template_row["account_name"])
        account_rows = []

        for row_index in range(template_row["row"], template_row["block_end_row"] + 1):
            invoice_number = worksheet.cell(row=row_index, column=4).value
            amount = worksheet.cell(row=row_index, column=5).value
            remark = worksheet.cell(row=row_index, column=7).value

            if invoice_number is None and amount is None and remark is None:
                continue

            account_rows.append(
                {
                    "invoice_number": str(invoice_number or "").strip(),
                    "amount": round(float(amount or 0), 2),
                    "remark": str(remark or "").strip(),
                }
            )

        if account_rows:
            existing_entries.setdefault(account_key, []).extend(account_rows)

    for account_key, entries in existing_entries.items():
        existing_entries[account_key] = deduplicate_entries(entries)

    return existing_entries


def merge_account_entries(existing_entries, new_entries):
    merged_entries = {}
    all_keys = set(existing_entries) | set(new_entries)

    for account_key in all_keys:
        combined_entries = []
        combined_entries.extend(existing_entries.get(account_key, []))
        combined_entries.extend(new_entries.get(account_key, []))

        if combined_entries:
            merged_entries[account_key] = deduplicate_entries(combined_entries)

    return merged_entries


def clear_row_values(worksheet, row_index):
    for column in ("C", "D", "E", "F", "G"):
        cell = worksheet[f"{column}{row_index}"]
        if not isinstance(cell, MergedCell):
            cell.value = None


def set_cell_value(worksheet, cell_ref, value):
    cell = worksheet[cell_ref]
    if not isinstance(cell, MergedCell):
        cell.value = value


def copy_row_style(worksheet, source_row, target_row):
    for column in range(1, worksheet.max_column + 1):
        source_cell = worksheet.cell(row=source_row, column=column)
        target_cell = worksheet.cell(row=target_row, column=column)
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)


def insert_entry_row(worksheet, source_row, insert_at):
    worksheet.insert_rows(insert_at)
    copy_row_style(worksheet, source_row, insert_at)


def reset_account_block(worksheet, start_row, end_row):
    if end_row > start_row:
        worksheet.delete_rows(start_row + 1, end_row - start_row)
    clear_row_values(worksheet, start_row)


def write_entry_to_row(
    worksheet,
    row_index,
    base_account,
    base_environment,
    billing_month_name,
    entry,
):
    set_cell_value(worksheet, f"A{row_index}", base_account)
    set_cell_value(worksheet, f"B{row_index}", base_environment)
    set_cell_value(worksheet, f"C{row_index}", billing_month_name)
    set_cell_value(worksheet, f"D{row_index}", entry["invoice_number"])
    set_cell_value(worksheet, f"E{row_index}", entry["amount"])
    set_cell_value(worksheet, f"G{row_index}", entry["remark"])


def write_template_row_entries(worksheet, template_row, entries, billing_month_name=None):
    row_index = template_row["row"]
    base_account = worksheet[f"A{row_index}"].value
    base_environment = worksheet[f"B{row_index}"].value
    base_month = billing_month_name or worksheet[f"C{row_index}"].value

    reset_account_block(worksheet, row_index, template_row["block_end_row"])
    if not entries:
        return

    write_entry_to_row(
        worksheet,
        row_index,
        base_account,
        base_environment,
        base_month,
        entries[0],
    )

    current_row = row_index
    for entry in entries[1:]:
        insert_at = current_row + 1
        insert_entry_row(worksheet, current_row, insert_at)
        current_row = insert_at
        clear_row_values(worksheet, current_row)
        write_entry_to_row(
            worksheet,
            current_row,
            base_account,
            base_environment,
            base_month,
            entry,
        )


def group_template_rows_by_sheet(template_rows):
    rows_by_sheet = {}
    for item in template_rows:
        rows_by_sheet.setdefault(item["sheet"], []).append(item)

    for sheet_rows in rows_by_sheet.values():
        sheet_rows.sort(key=lambda item: item["row"], reverse=True)

    return rows_by_sheet


def prepare_worksheet_for_writing(worksheet):
    ranges_to_unmerge = [str(cell_range) for cell_range in worksheet.merged_cells.ranges]
    for cell_range in ranges_to_unmerge:
        worksheet.unmerge_cells(cell_range)


def rebuild_title_merges(worksheet):
    for row_index in range(1, worksheet.max_row + 1):
        account_value = worksheet.cell(row=row_index, column=1).value
        if not str(account_value or "").strip():
            continue

        if not normalize_text(account_value).startswith("invoice"):
            continue

        worksheet.merge_cells(
            start_row=row_index,
            start_column=1,
            end_row=row_index,
            end_column=7,
        )


def rebuild_account_merges(worksheet):
    row_index = 1
    while row_index <= worksheet.max_row:
        account_value = worksheet.cell(row=row_index, column=1).value
        environment_value = worksheet.cell(row=row_index, column=2).value

        if not row_has_account_template_values(account_value, environment_value):
            row_index += 1
            continue

        group_account = account_value
        group_environment = environment_value
        start_row = row_index
        end_row = row_index
        scan_row = row_index + 1

        while scan_row <= worksheet.max_row:
            next_account = worksheet.cell(row=scan_row, column=1).value
            next_environment = worksheet.cell(row=scan_row, column=2).value

            if next_account != group_account or next_environment != group_environment:
                break

            end_row = scan_row
            scan_row += 1

        if end_row > start_row:
            worksheet.merge_cells(
                start_row=start_row,
                start_column=1,
                end_row=end_row,
                end_column=1,
            )
            worksheet.merge_cells(
                start_row=start_row,
                start_column=2,
                end_row=end_row,
                end_column=2,
            )

        row_index = scan_row


def rebuild_section_totals(worksheet):
    row_index = 1
    while row_index <= worksheet.max_row:
        account_value = worksheet.cell(row=row_index, column=1).value

        if normalize_text(account_value) != "account":
            row_index += 1
            continue

        start_row = row_index + 1
        scan_row = start_row
        end_row = None

        while scan_row <= worksheet.max_row:
            current_account = worksheet.cell(row=scan_row, column=1).value
            current_month = worksheet.cell(row=scan_row, column=3).value
            current_invoice = worksheet.cell(row=scan_row, column=4).value
            current_amount = worksheet.cell(row=scan_row, column=5).value

            if is_title_or_header(current_account):
                break

            if (
                current_account is None
                and current_month is None
                and current_invoice is None
                and current_amount is None
            ):
                break

            if current_amount is not None or current_invoice is not None or current_month is not None:
                end_row = scan_row

            scan_row += 1

        if end_row is not None:
            set_cell_value(worksheet, f"F{start_row}", f"=SUM(E{start_row}:E{end_row})")
            for clear_row in range(start_row + 1, end_row + 1):
                set_cell_value(worksheet, f"F{clear_row}", None)

            if end_row > start_row:
                worksheet.merge_cells(
                    start_row=start_row,
                    start_column=6,
                    end_row=end_row,
                    end_column=6,
                )

        row_index = scan_row


def build_account_name_lookup(data):
    account_names = {}
    for account in data:
        account_key = normalize_account_lookup(account["account_name"])
        account_names.setdefault(account_key, str(account["account_name"]).strip())
    return account_names


def style_generated_workbook(workbook):
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    title_fill = PatternFill("solid", fgColor="1F4E78")
    thin_side = Side(style="thin", color="A6A6A6")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A3"
        worksheet.column_dimensions["A"].width = 28
        worksheet.column_dimensions["B"].width = 18
        worksheet.column_dimensions["C"].width = 14
        worksheet.column_dimensions["D"].width = 22
        worksheet.column_dimensions["E"].width = 14
        worksheet.column_dimensions["F"].width = 16
        worksheet.column_dimensions["G"].width = 55

        for row in worksheet.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for cell in worksheet[1]:
            cell.fill = title_fill
            cell.font = Font(color="FFFFFF", bold=True, size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for cell in worksheet[2]:
            cell.fill = header_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_index in range(3, worksheet.max_row + 1):
            worksheet.cell(row=row_index, column=5).number_format = '$#,##0.00'
            worksheet.cell(row=row_index, column=6).number_format = '$#,##0.00'
            if worksheet.cell(row=row_index, column=1).value == "Grand Total":
                for column_index in range(1, 8):
                    cell = worksheet.cell(row=row_index, column=column_index)
                    cell.fill = header_fill
                    cell.font = Font(bold=True)


def export_to_generated_excel(data, output_file, billing_month_name=None):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoice Report"

    headers = [
        "Account",
        "Environment Type",
        "Month",
        "Invoice No",
        "Amount",
        "Total",
        "Remarks",
    ]
    worksheet.append([f"Invoice Report - {billing_month_name or ''}".strip(), None, None, None, None, None, None])
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    worksheet.append(headers)

    account_entries = build_account_entries(data)
    account_names = build_account_name_lookup(data)
    current_row = 3
    account_total_cells = []

    for account_key in sorted(account_entries, key=lambda key: account_names.get(key, key).lower()):
        entries = account_entries[account_key]
        if not entries:
            continue

        account_name = account_names.get(account_key, account_key)
        start_row = current_row
        for entry in entries:
            worksheet.append(
                [
                    account_name,
                    "AWS",
                    billing_month_name,
                    entry["invoice_number"],
                    entry["amount"],
                    None,
                    entry["remark"],
                ]
            )
            current_row += 1

        end_row = current_row - 1
        worksheet.cell(row=start_row, column=6).value = f"=SUM(E{start_row}:E{end_row})"
        account_total_cells.append(f"F{start_row}")
        if end_row > start_row:
            worksheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
            worksheet.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
            worksheet.merge_cells(start_row=start_row, start_column=6, end_row=end_row, end_column=6)

        worksheet.append([None, None, None, None, None, None, None])
        current_row += 1

    if account_total_cells:
        total_row = current_row
        worksheet.append(["Grand Total", None, None, None, None, f"=SUM({','.join(account_total_cells)})", None])
        worksheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)

    style_generated_workbook(workbook)

    try:
        workbook.save(output_file)
    except PermissionError:
        raise PermissionError(
            f"Unable to update '{output_file}'. Please close the Excel file and run again."
        )

    return output_file


def export_to_formatted_excel(data, template_file, output_file, billing_month_name=None):
    if not template_file:
        return export_to_generated_excel(data, output_file, billing_month_name)

    workbook = load_workbook(template_file)
    for worksheet in workbook.worksheets:
        prepare_worksheet_for_writing(worksheet)

    # Treat the workbook as a layout template and rebuild account rows from the
    # current invoice data so reruns do not carry forward old values.
    pending_account_entries = build_account_entries(data)
    template_rows = collect_template_rows(workbook)
    rows_by_sheet = group_template_rows_by_sheet(template_rows)

    for sheet_name, sheet_rows in rows_by_sheet.items():
        worksheet = workbook[sheet_name]
        for template_row in sheet_rows:
            account_key = normalize_account_lookup(template_row["account_name"])
            entries = pending_account_entries.pop(account_key, [])
            write_template_row_entries(worksheet, template_row, entries, billing_month_name)

    for worksheet in workbook.worksheets:
        rebuild_title_merges(worksheet)
        rebuild_section_totals(worksheet)
        rebuild_account_merges(worksheet)

    try:
        workbook.save(output_file)
    except PermissionError:
        raise PermissionError(
            f"Unable to update '{output_file}'. Please close the Excel file and run again."
        )

    return output_file


if __name__ == "__main__":
    print("Validating invoice billing period...")
    billing_month_name, billing_period_label = validate_invoices_match_previous_month(PDF_DIR)

    output_file = build_output_file_name(billing_period_label)
    print(f"Processing invoices for: {billing_period_label}")

    print("Processing invoices...")
    data = process_all_invoices(PDF_DIR, billing_period_label)

    print("Exporting to Excel...")
    saved_output_file = export_to_generated_excel(data, output_file, billing_month_name)

    print(
        f"Invoice Report is updated for {billing_period_label} : "
        f"{os.path.basename(saved_output_file)}"
    )
