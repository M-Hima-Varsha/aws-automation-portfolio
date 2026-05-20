from __future__ import annotations

import csv
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import date

import boto3
from botocore.exceptions import BotoCoreError, ClientError, ProfileNotFound
from dateutil.relativedelta import relativedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


WORK_DIR = os.path.dirname(os.path.abspath(__file__))


@dataclass(frozen=True)
class AccountConfig:
    name: str
    account_id: str


PRICING_SERVICE_CODES = {
    "Amazon ElastiCache": "AmazonElastiCache",
    "Amazon Relational Database Service": "AmazonRDS",
    "Amazon Redshift": "AmazonRedshift",
    "Amazon OpenSearch Service": "AmazonOpenSearch",
}

RESERVATION_SERVICE_FALLBACKS = [
    "Amazon Elastic Compute Cloud - Compute",
    "Amazon Relational Database Service",
    "Amazon ElastiCache",
    "Amazon Redshift",
    "Amazon Elasticsearch Service",
    "Amazon OpenSearch Service",
]

DETAIL_CSV_HEADERS = [
    "Service",
    "Account name",
    "Subscription ID",
    "Reservation ID",
    "Instance type",
    "Utilization",
    "Reservation hours purchased",
    "Reservation hours used",
    "Reservation hours unused",
    "Used units",
    "Account ID",
    "Start date",
    "End date",
    "Reserved count",
    "Region",
    "Availability Zone",
    "Product description",
    "Scope",
    "Tenancy",
    "Payment option",
    "Offering type",
    "On-Demand cost equivalent",
    "Amortized upfront fee",
    "Amortized recurring fee",
    "Effective reservation cost",
    "Net savings",
    "Potential savings",
    "Average On-Demand hourly rate",
    "Total asset value",
    "Effective hourly rate",
    "Upfront fee",
    "Hourly recurring fee",
    "Cost for unused hours",
]

SERVICE_BLOCK_LAYOUTS = {
    "Amazon Relational Database Service": {
        "title": "RDS RI Utilization",
        "headers": [
            "Account",
            "Subscription",
            "RI PURCHASED",
            "RI USED",
            "PURCHASED",
            "vCPU",
            "MEMORY",
            "NF",
            "RI HOURS",
            "RI HOURS Used",
            "RI UTILIZATION %",
            "TOTAL RI UTILIZATION%",
        ],
    },
    "Amazon ElastiCache": {
        "title": "ElastiCache RI Utilization (Valkey)",
        "headers": [
            "Account",
            "Subscription",
            "Node Tyoe",
            "Purchased",
            "vCPU",
            "MEMORY",
            "RI HOURS",
            "RI HOURS Used",
            "RI HOURS Unused",
            "Number of Reserved Nodes",
            "Cache Utilization %",
            "Total Cache Utilization %",
        ],
    },
}

GENERIC_HEADERS = [
    "Account",
    "Service",
    "Subscription",
    "Instance Type",
    "Purchased",
    "Used",
    "vCPU",
    "MEMORY",
    "NF",
    "RI HOURS",
    "RI HOURS Used",
    "RI HOURS Unused",
    "RI UTILIZATION %",
    "Region",
    "Scope",
    "Product Description",
]

DOWNLOAD_FILE_PREFIX = "reservations-utilization-table"


def get_reporting_period(today: date | None = None) -> tuple[date, date, str, str]:
    today = today or date.today()
    previous_month_start = today.replace(day=1) - relativedelta(months=1)
    current_month_start = today.replace(day=1)
    return (
        previous_month_start,
        current_month_start,
        previous_month_start.strftime("%B"),
        previous_month_start.strftime("%Y"),
    )


def print_user_error(message: str) -> None:
    sys.stderr.write(f"Error: {message}\n")


def build_output_paths(month_name: str, year: str) -> tuple[str, str]:
    report_folder = os.path.join(WORK_DIR, f"{month_name} - {year}")
    output_excel = os.path.join(WORK_DIR, f"RI Utilization Report {month_name} - {year}.xlsx")
    return report_folder, output_excel


def to_float(value: object) -> float:
    if value in (None, ""):
        return 0.0
    text = str(value).replace(",", "").replace("%", "").strip()
    match = re.search(r"-?\d+(\.\d+)?", text)
    return float(match.group()) if match else 0.0


def format_number(value: float) -> int | float:
    rounded = round(value, 2)
    return int(rounded) if rounded.is_integer() else rounded


def format_percent(value: float) -> str:
    return f"{round(value, 2)}%"


def format_decimal(value: object) -> str:
    number = to_float(value)
    if number.is_integer():
        return str(int(number))
    return str(round(number, 10)).rstrip("0").rstrip(".")


def parse_number(value: object) -> int | float | str:
    text = str(value).strip()
    if not text:
        return ""
    number = to_float(text)
    if str(number) == "0.0" and text not in {"0", "0.0", "0.00"}:
        return text
    return int(number) if number.is_integer() else number


def format_utilization_for_sheet(value: object) -> str:
    text = str(value).strip()
    if not text:
        return ""

    if "%" in text:
        return format_percent(to_float(text))

    number = to_float(text)
    if number <= 1:
        number *= 100
    return format_percent(number)


def get_selected_profile() -> str | None:
    return os.environ.get("AWS_PROFILE") or os.environ.get("AWS_DEFAULT_PROFILE")


def create_session() -> boto3.session.Session:
    profile_name = get_selected_profile()
    if profile_name:
        return boto3.Session(profile_name=profile_name)
    return boto3.Session()


def build_service_filter(account_id: str, service_name: str) -> dict:
    return {
        "And": [
            {"Dimensions": {"Key": "LINKED_ACCOUNT", "Values": [account_id]}},
            {"Dimensions": {"Key": "SERVICE", "Values": [service_name]}},
        ]
    }


def build_account_filter(account_id: str) -> dict:
    return {"Dimensions": {"Key": "LINKED_ACCOUNT", "Values": [account_id]}}


def paginate_reservation_utilization(client, **kwargs) -> list[dict]:
    pages: list[dict] = []
    next_token: str | None = None
    while True:
        request = dict(kwargs)
        if next_token:
            request["NextPageToken"] = next_token
        response = client.get_reservation_utilization(**request)
        pages.append(response)
        next_token = response.get("NextPageToken")
        if not next_token:
            return pages


def fetch_service_names(client, time_period: dict[str, str], account: AccountConfig) -> list[str]:
    service_names: list[str] = []

    next_page_token: str | None = None
    while True:
        request = {
            "TimePeriod": time_period,
            "Context": "RESERVATIONS",
            "Dimension": "SERVICE",
            "Filter": build_account_filter(account.account_id),
        }
        if next_page_token:
            request["NextPageToken"] = next_page_token
        try:
            response = client.get_dimension_values(**request)
        except ClientError:
            service_names = []
            break

        for value in response.get("DimensionValues", []):
            service_name = str(value.get("Value", "")).strip()
            if service_name and service_name not in service_names:
                service_names.append(service_name)

        next_page_token = response.get("NextPageToken")
        if not next_page_token:
            break

    if service_names:
        return service_names

    for service_name in RESERVATION_SERVICE_FALLBACKS:
        totals = fetch_total_values(client, time_period, account, service_name)
        if (
            to_float(totals.get("PurchasedHours")) == 0
            and to_float(totals.get("TotalActualHours")) == 0
            and to_float(totals.get("UnusedHours")) == 0
        ):
            continue
        service_names.append(service_name)
    return service_names


def fetch_group_values(
    client,
    time_period: dict[str, str],
    account: AccountConfig,
    service_name: str,
    group_key: str,
) -> list[dict]:
    try:
        pages = paginate_reservation_utilization(
            client,
            TimePeriod=time_period,
            Filter=build_service_filter(account.account_id, service_name),
            GroupBy=[{"Type": "DIMENSION", "Key": group_key}],
        )
    except ClientError as exc:
        error = exc.response.get("Error", {})
        if error.get("Code") == "ValidationException" and group_key == "INSTANCE_TYPE":
            return []
        raise
    groups: list[dict] = []
    for page in pages:
        for block in page.get("UtilizationsByTime", []):
            groups.extend(block.get("Groups", []))
    return groups


def fetch_total_values(
    client,
    time_period: dict[str, str],
    account: AccountConfig,
    service_name: str,
) -> dict[str, str]:
    pages = paginate_reservation_utilization(
        client,
        TimePeriod=time_period,
        Granularity="MONTHLY",
        Filter=build_service_filter(account.account_id, service_name),
    )
    totals: dict[str, str] = {}
    for page in pages:
        total = page.get("Total") or {}
        if total:
            totals = total
    return totals


def get_subscription_text(groups: list[dict]) -> str:
    values = [
        str(group.get("Value") or group.get("Key") or "").strip()
        for group in groups
        if group.get("Value") or group.get("Key")
    ]
    unique_values: list[str] = []
    for value in values:
        if value not in unique_values:
            unique_values.append(value)
    return ", ".join(unique_values[:3])


def get_primary_group(groups: list[dict]) -> dict:
    return max(
        groups,
        key=lambda group: to_float(group.get("Utilization", {}).get("PurchasedHours")),
        default={},
    )


def get_primary_instance_type(groups: list[dict]) -> str:
    attributes = get_primary_group(groups).get("Attributes", {})
    return get_instance_type(attributes)


def get_instance_type(attributes: dict[str, object]) -> str:
    return str(
        attributes.get("instanceType")
        or attributes.get("cacheNodeType")
        or attributes.get("instanceFamily")
        or ""
    ).strip()


def get_normalization_factor(instance_type: str) -> int | float | str:
    if not instance_type:
        return ""
    size = instance_type.split(".")[-1].lower()
    if size in {"large", "xlarge"}:
        return 1
    match = re.fullmatch(r"(\d+)xlarge", size)
    if match:
        return int(match.group(1))
    return ""


def get_pricing_specs(session: boto3.session.Session, service_name: str, instance_type: str) -> tuple[str, str]:
    if not instance_type:
        return "", ""

    service_code = PRICING_SERVICE_CODES.get(service_name)
    if not service_code:
        return "", ""

    try:
        pricing = session.client("pricing")
        response = pricing.get_products(
            ServiceCode=service_code,
            Filters=[
                {"Type": "TERM_MATCH", "Field": "instanceType", "Value": instance_type},
            ],
            MaxResults=20,
        )
    except (BotoCoreError, ClientError):
        return "", ""

    for price_item in response.get("PriceList", []):
        item = json.loads(price_item)
        attributes = item.get("product", {}).get("attributes", {})
        if attributes.get("instanceType") != instance_type:
            continue
        vcpu = attributes.get("vcpu") or attributes.get("vcpuPerInstance") or ""
        memory = attributes.get("memory") or ""
        if vcpu or memory:
            return str(vcpu), str(memory)

    return "", ""


def write_csv_snapshot(
    csv_path: str,
    rows: list[dict[str, object]],
) -> None:
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as file_handle:
        writer = csv.DictWriter(file_handle, fieldnames=DETAIL_CSV_HEADERS, quoting=csv.QUOTE_ALL)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def format_region_name(region_code: str) -> str:
    region_names = {
        "us-east-1": "US East (N. Virginia)",
        "us-east-2": "US East (Ohio)",
        "us-west-1": "US West (N. California)",
        "us-west-2": "US West (Oregon)",
    }
    return region_names.get(region_code, region_code)


def sanitize_service_name(service_name: str) -> str:
    sanitized = re.sub(r"[^A-Za-z0-9._ -]+", "", service_name).strip()
    return sanitized or "UnknownService"


def build_detail_csv_rows(service_name: str, groups: list[dict]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for group in groups:
        attributes = group.get("Attributes", {})
        utilization = group.get("Utilization", {})
        row = {
            "Service": service_name,
            "Account name": str(attributes.get("accountName", "")).strip(),
            "Subscription ID": str(group.get("Value") or attributes.get("subscriptionId", "")).strip(),
            "Reservation ID": str(attributes.get("leaseId", "")).strip(),
            "Instance type": get_instance_type(attributes),
            "Utilization": format_decimal(to_float(utilization.get("UtilizationPercentage")) / 100),
            "Reservation hours purchased": format_decimal(utilization.get("PurchasedHours")),
            "Reservation hours used": format_decimal(utilization.get("TotalActualHours")),
            "Reservation hours unused": format_decimal(utilization.get("UnusedHours")),
            "Used units": format_decimal(utilization.get("TotalActualUnits")),
            "Account ID": str(attributes.get("accountId", "")).strip(),
            "Start date": str(attributes.get("startDateTime", "")).strip(),
            "End date": str(attributes.get("endDateTime", "")).strip(),
            "Reserved count": format_decimal(attributes.get("numberOfInstances")),
            "Region": format_region_name(str(attributes.get("region", "")).strip()),
            "Availability Zone": str(attributes.get("availabilityZone", "")).strip(),
            "Product description": str(attributes.get("platform", "")).strip(),
            "Scope": str(attributes.get("scope", "")).strip(),
            "Tenancy": str(attributes.get("tenancy", "")).strip(),
            "Payment option": str(attributes.get("subscriptionType", "")).strip(),
            "Offering type": str(attributes.get("offeringType", "")).strip(),
            "On-Demand cost equivalent": format_decimal(utilization.get("OnDemandCostOfRIHoursUsed")),
            "Amortized upfront fee": format_decimal(utilization.get("AmortizedUpfrontFee")),
            "Amortized recurring fee": format_decimal(utilization.get("AmortizedRecurringFee")),
            "Effective reservation cost": format_decimal(utilization.get("TotalAmortizedFee")),
            "Net savings": format_decimal(utilization.get("NetRISavings")),
            "Potential savings": format_decimal(utilization.get("TotalPotentialRISavings")),
            "Average On-Demand hourly rate": format_decimal(attributes.get("averageOnDemandHourlyRate")),
            "Total asset value": format_decimal(attributes.get("totalAssetValue")),
            "Effective hourly rate": format_decimal(attributes.get("effectiveHourlyRate")),
            "Upfront fee": format_decimal(attributes.get("upfrontFee")),
            "Hourly recurring fee": format_decimal(attributes.get("hourlyRecurringFee")),
            "Cost for unused hours": format_decimal(utilization.get("RICostForUnusedHours")),
        }
        rows.append(row)
    return rows


def get_account_display_name(session: boto3.session.Session, account_id: str) -> str:
    try:
        organizations = session.client("organizations")
        account = organizations.describe_account(AccountId=account_id).get("Account", {})
        account_name = str(account.get("Name", "")).strip()
        if account_name:
            return account_name
    except (BotoCoreError, ClientError):
        pass
    return account_id


def fetch_service_data(
    session: boto3.session.Session,
    account: AccountConfig,
    service_name: str,
    period_start: date,
    period_end: date,
) -> dict[str, object] | None:
    client = session.client("ce")
    time_period = {
        "Start": period_start.isoformat(),
        "End": period_end.isoformat(),
    }

    totals = fetch_total_values(client, time_period, account, service_name)
    purchased_units = to_float(totals.get("PurchasedUnits"))
    purchased_hours = to_float(totals.get("PurchasedHours"))
    used_units = to_float(totals.get("TotalActualUnits"))
    used_hours = to_float(totals.get("TotalActualHours"))
    unused_hours = to_float(totals.get("UnusedHours"))

    if purchased_units == 0 and purchased_hours == 0 and used_units == 0 and used_hours == 0:
        return None

    subscription_groups = fetch_group_values(client, time_period, account, service_name, "SUBSCRIPTION_ID")

    subscription_text = get_subscription_text(subscription_groups)
    instance_type = get_primary_instance_type(subscription_groups)
    vcpu, memory = get_pricing_specs(session, service_name, instance_type)
    normalization_factor = get_normalization_factor(instance_type)

    utilization_pct = to_float(totals.get("UtilizationPercentage") or 0)

    result: dict[str, object] = {
        "service": service_name,
        "subscription": subscription_text,
        "purchased": format_number(purchased_units),
        "vcpu": vcpu,
        "memory": memory,
        "nf": normalization_factor,
        "ri_hours": format_number(purchased_hours),
        "ri_hours_used": format_number(used_hours),
        "ri_hours_unused": format_number(unused_hours),
        "ri_utilization_pct": format_percent(utilization_pct),
        "instance_type": instance_type,
        "detail_csv_rows": build_detail_csv_rows(service_name, subscription_groups),
    }

    return result


def load_detail_csv_rows(report_folder: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for file_name in sorted(os.listdir(report_folder)):
        if not file_name.startswith(DOWNLOAD_FILE_PREFIX):
            continue
        csv_path = os.path.join(report_folder, file_name)
        with open(csv_path, newline="", encoding="utf-8-sig") as file_handle:
            reader = csv.DictReader(file_handle)
            rows.extend(dict(row) for row in reader)
    return rows


def get_prepared_account_name(report_folder: str, account_id: str, fallback_name: str) -> str:
    for row in load_detail_csv_rows(report_folder):
        row_account_id = str(row.get("Account ID", "")).strip()
        row_account_name = str(row.get("Account name", "")).strip()
        if row_account_id == account_id and row_account_name:
            return row_account_name
    return fallback_name


def build_workbook_row(
    session: boto3.session.Session,
    csv_row: dict[str, str],
    ri_hours_formula: str,
) -> dict[str, object]:
    instance_type = csv_row.get("Instance type", "")
    service_name = csv_row.get("Service", "")
    vcpu, memory = get_pricing_specs(session, service_name, instance_type)
    utilization = format_utilization_for_sheet(csv_row.get("Utilization", ""))
    return {
        "Account": csv_row.get("Account name", ""),
        "Service": service_name,
        "Subscription": parse_number(csv_row.get("Subscription ID", "")),
        "Instance Type": instance_type,
        "Purchased": parse_number(csv_row.get("Reserved count", "")),
        "Used": parse_number(csv_row.get("Used units", "")),
        "vCPU": parse_number(vcpu),
        "MEMORY": parse_number(memory),
        "RI HOURS": ri_hours_formula,
        "RI HOURS Used": parse_number(csv_row.get("Reservation hours used", "")),
        "RI HOURS Unused": parse_number(csv_row.get("Reservation hours unused", "")),
        "NF": get_normalization_factor(instance_type),
        "RI UTILIZATION %": utilization,
        "Region": csv_row.get("Region", ""),
        "Scope": csv_row.get("Scope", ""),
        "Product Description": csv_row.get("Product description", ""),
    }


def build_service_block_row(service_name: str, row: dict[str, object]) -> dict[str, object]:
    if service_name == "Amazon Relational Database Service":
        return {
            "Account": row.get("Account", ""),
            "Subscription": row.get("Subscription", ""),
            "RI PURCHASED": row.get("Instance Type", ""),
            "RI USED": row.get("Instance Type", ""),
            "PURCHASED": row.get("Purchased", ""),
            "vCPU": row.get("vCPU", ""),
            "MEMORY": row.get("MEMORY", ""),
            "NF": row.get("NF", ""),
            "RI HOURS": row.get("RI HOURS", ""),
            "RI HOURS Used": row.get("RI HOURS Used", ""),
            "RI UTILIZATION %": row.get("RI UTILIZATION %", ""),
            "TOTAL RI UTILIZATION%": row.get("RI UTILIZATION %", ""),
        }

    if service_name == "Amazon ElastiCache":
        return {
            "Account": row.get("Account", ""),
            "Subscription": row.get("Subscription", ""),
            "Node Tyoe": row.get("Instance Type", ""),
            "Purchased": row.get("Purchased", ""),
            "vCPU": row.get("vCPU", ""),
            "MEMORY": row.get("MEMORY", ""),
            "RI HOURS": row.get("RI HOURS", ""),
            "RI HOURS Used": row.get("RI HOURS Used", ""),
            "RI HOURS Unused": row.get("RI HOURS Unused", ""),
            "Number of Reserved Nodes": row.get("Purchased", ""),
            "Cache Utilization %": row.get("RI UTILIZATION %", ""),
            "Total Cache Utilization %": row.get("RI UTILIZATION %", ""),
        }

    return row


def style_service_block(
    sheet,
    title: str,
    title_row: int,
    header_row: int,
    data_start_row: int,
    data_end_row: int,
    header_count: int,
) -> None:
    title_fill = PatternFill(fill_type="solid", fgColor="BFBFBF")
    header_fill = PatternFill(fill_type="solid", fgColor="D9D9D9")
    data_fill = PatternFill(fill_type="solid", fgColor="F2F2F2")
    bold_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_index in range(1, header_count + 1):
        cell = sheet.cell(title_row, col_index)
        cell.fill = title_fill
        cell.font = bold_font
        cell.alignment = center
        cell.border = border
    sheet.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=header_count)
    sheet.cell(title_row, 1).value = title

    for col_index in range(1, header_count + 1):
        cell = sheet.cell(header_row, col_index)
        cell.fill = header_fill
        cell.font = bold_font
        cell.alignment = center
        cell.border = border

    for row_index in range(data_start_row, data_end_row + 1):
        for col_index in range(1, header_count + 1):
            cell = sheet.cell(row_index, col_index)
            cell.fill = data_fill
            cell.alignment = center
            cell.border = border


def create_output_workbook(
    output_path: str,
    session: boto3.session.Session,
    report_folder: str,
    period_start: date,
    period_end: date,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RI Utilization"

    ri_hours_formula = f"={(period_end - period_start).days}*24"
    workbook_rows = [
        build_workbook_row(session, row, ri_hours_formula)
        for row in load_detail_csv_rows(report_folder)
    ]

    preferred_service_order = [
        "Amazon Relational Database Service",
        "Amazon ElastiCache",
    ]
    grouped_rows: dict[str, list[dict[str, object]]] = {}
    for row in workbook_rows:
        service_name = str(row.get("Service", "")).strip() or "Other"
        grouped_rows.setdefault(service_name, []).append(row)

    ordered_service_names = [
        service_name for service_name in preferred_service_order if service_name in grouped_rows
    ]
    ordered_service_names.extend(
        sorted(service_name for service_name in grouped_rows if service_name not in ordered_service_names)
    )

    current_row = 1
    for service_name in ordered_service_names:
        layout = SERVICE_BLOCK_LAYOUTS.get(
            service_name,
            {"title": f"{service_name} RI Utilization", "headers": GENERIC_HEADERS},
        )
        headers = layout["headers"]
        title = layout["title"]
        sheet.cell(current_row, 1).value = title
        title_row = current_row
        current_row += 1

        for col_index, header in enumerate(headers, start=1):
            sheet.cell(current_row, col_index).value = header
        header_row = current_row
        current_row += 1

        data_start_row = current_row
        for row_values in grouped_rows[service_name]:
            block_row = build_service_block_row(service_name, row_values)
            for col_index, header in enumerate(headers, start=1):
                sheet.cell(current_row, col_index).value = block_row.get(header, "")
            current_row += 1
        data_end_row = current_row - 1

        style_service_block(
            sheet,
            title,
            title_row,
            header_row,
            data_start_row,
            data_end_row,
            len(headers),
        )

        current_row += 2

    column_widths = {
        "A": 18,
        "B": 15,
        "C": 17,
        "D": 17,
        "E": 12,
        "F": 15,
        "G": 18,
        "H": 20,
        "I": 16,
        "J": 16,
        "K": 18,
        "L": 22,
        "M": 18,
        "N": 16,
        "O": 16,
        "P": 22,
    }
    for column_letter, width in column_widths.items():
        sheet.column_dimensions[column_letter].width = width

    workbook.save(output_path)

def main() -> int:
    period_start, period_end, month_name, year = get_reporting_period()
    output_dir, output_path = build_output_paths(month_name, year)
    os.makedirs(output_dir, exist_ok=True)

    try:
        session = create_session()
        sts = session.client("sts")
        caller = sts.get_caller_identity()
        account_id = str(caller.get("Account", "")).strip()
        account_name = get_account_display_name(session, account_id)
    except ProfileNotFound as exc:
        print_user_error(f"AWS profile not found: {exc}")
        return 1
    except (BotoCoreError, ClientError) as exc:
        print_user_error(f"Unable to use AWS credentials from .aws files: {exc}")
        return 1

    if not account_id:
        print_user_error("Unable to determine the AWS account from the configured credentials.")
        return 1

    account_results: dict[tuple[str, str], dict[str, object]] = {}
    current_account = AccountConfig(name=account_name, account_id=account_id)

    try:
        ce_client = session.client("ce")
        time_period = {
            "Start": period_start.isoformat(),
            "End": period_end.isoformat(),
        }
        service_names = fetch_service_names(ce_client, time_period, current_account)
        for service_name in service_names:
            values = fetch_service_data(session, current_account, service_name, period_start, period_end)
            if not values:
                continue

            account_results[(current_account.name, service_name)] = values
            csv_name = (
                f"reservations-utilization-table - "
                f"{current_account.name} - {sanitize_service_name(service_name)}.csv"
            )
            write_csv_snapshot(
                os.path.join(output_dir, csv_name),
                values.get("detail_csv_rows", []),
            )

    except (BotoCoreError, ClientError) as exc:
        print_user_error(f"Failed to fetch RI utilization data from AWS: {exc}")
        return 1

    if not account_results:
        print_user_error("No RI utilization data was returned for the configured accounts.")
        return 1

    create_output_workbook(output_path, session, output_dir, period_start, period_end)
    prepared_account_name = get_prepared_account_name(
        output_dir,
        current_account.account_id,
        current_account.name,
    )
    print(
        f"RI Utilization report is updated for {prepared_account_name} : "
        f"{os.path.basename(output_path)}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
